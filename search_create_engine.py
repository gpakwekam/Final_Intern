import streamlit as st
import fitz  # PyMuPDF for PDFs
import docx  # python-docx for Word docs
import openpyxl  # openpyxl for Excel
from openpyxl import Workbook
import re  # Regex for sentence splitting

def extract_sentences(text, search_term):
    # Split text into sentences
    sentences = re.split(r'(?<!\w\.\w.)(?<![A-Z][a-z]\.)(?<=\.|\?)\s', text)
    # Filter sentences that contain the search term
    return [sentence.strip() for sentence in sentences if search_term.lower() in sentence.lower()]

def extract_from_word(doc_path, search_term):
    doc = docx.Document(doc_path)
    results = []

    for para in doc.paragraphs:
        if search_term.lower() in para.text.lower():  # Check if the paragraph contains the search term
            matching_sentences = extract_sentences(para.text, search_term)
            results.extend([(sentence,) for sentence in matching_sentences])

    return results

def extract_from_pdf(pdf_path, search_term):
    doc = fitz.open(pdf_path)
    results = []

    for page_num in range(len(doc)):
        page = doc.load_page(page_num)
        text = page.get_text("text")
        if search_term.lower() in text.lower():  # Check if the page contains the search term
            matching_sentences = extract_sentences(text, search_term)
            results.extend([(page_num + 1, sentence) for sentence in matching_sentences])

    return results

def save_to_excel(data, output_path, headers):
    wb = Workbook()
    ws = wb.active
    ws.append(headers)

    for item in data:
        ws.append(item)

    wb.save(output_path)

def main():
    st.title("Search and Create Engine")
    st.write("Upload a document, define the search term, and get an Excel report.")

    uploaded_file = st.file_uploader("Upload a file", type=["pdf", "docx"])
    search_term = st.text_input("Enter the term to search for")
    output_param = st.text_input("Enter the parameter for the Excel outcome")

    if st.button("Process"):
        if uploaded_file and search_term and output_param:
            file_path = f"temp_{uploaded_file.name}"
            with open(file_path, "wb") as f:
                f.write(uploaded_file.getbuffer())
            
            if file_path.endswith(".docx"):
                results = extract_from_word(file_path, search_term)
                headers = [output_param]
                save_to_excel(results, "output.xlsx", headers)

            elif file_path.endswith(".pdf"):
                results = extract_from_pdf(file_path, search_term)
                headers = ["Page Number", output_param]
                save_to_excel(results, "output.xlsx", headers)

            st.success("Processing complete. Download your file below.")
            with open("output.xlsx", "rb") as file:
                st.download_button(label="Download Excel file", data=file, file_name="output.xlsx", mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        else:
            st.error("Please upload a file, enter a search term, and specify an output parameter.")

if __name__ == "__main__":
    main()
